from __future__ import annotations

import json

import openpyxl
import pandas as pd

from itgc_uar.cli import main
from itgc_uar.reporting.dashboard import build_payload


def test_cli_all_writes_every_output_and_is_consistent(tmp_path):
    data, out = tmp_path / "d", tmp_path / "o"
    assert main(["all", "--profile", "baseline", "--data", str(data), "--out", str(out)]) == 0
    for f in ("ITGC_UAR_Exception_Report.xlsx", "ITGC_UAR_Dashboard.html", "run_manifest.json", "run.log",
              "csv/exception_register.csv", "csv/control_summary.csv", "csv/T6_exceptions.csv"):
        assert (out / f).is_file(), f
    reg = pd.read_csv(out / "csv/exception_register.csv")
    assert len(reg) == 82 and (reg["severity"] == "High").sum() == 29
    wb = openpyxl.load_workbook(out / "ITGC_UAR_Exception_Report.xlsx")
    assert wb.sheetnames[:4] == ["Cover", "Summary", "Exception Register", "Remediation Plan"]
    assert {"T1", "T2", "T3", "T4", "T5", "T6", "Reconciliation", "Data Quality", "Synthetic Validation"} <= set(wb.sheetnames)
    assert wb["Exception Register"].max_row - 1 == 82
    assert "SIMULATED" in wb["Cover"]["B2"].value
    assert sum(wb[c].max_row - 1 for c in ("T1", "T2", "T3", "T4", "T5", "T6")) == 82
    html = (out / "ITGC_UAR_Dashboard.html").read_text()
    assert "SIMULATED DATA" in html and "/*__DATA__*/null" not in html and "http://" not in html.replace("http://www.w3.org/2000/svg", "")
    m = json.loads((out / "run_manifest.json").read_text())
    assert m["total_exceptions"] == 82 and m["reconciled"] and m["data_is_simulated"]
    assert m["inputs"]["hr_master"]["sha256"].startswith("46ee7fce158f")


def test_dashboard_payload_matches_register(baseline_result):
    p = build_payload(baseline_result, {"profile": "baseline"}, None)
    assert len(p["register"]) == baseline_result.total_exceptions
    assert sum(r["exceptions"] for r in p["summary"]) == baseline_result.total_exceptions


def test_cli_exit_codes(tmp_path, capsys):
    assert main(["run", "--data", str(tmp_path / "nope"), "--out", str(tmp_path / "o")]) == 1
    data = tmp_path / "d"
    main(["generate", "--out", str(data)])
    bad = pd.read_csv(data / "user_access.csv", dtype=str)
    bad.loc[0, "account_status"] = "Locked"
    bad.to_csv(data / "user_access.csv", index=False)
    assert main(["run", "--data", str(data), "--out", str(tmp_path / "o")]) == 3
    assert main(["run", "--data", str(data), "--out", str(tmp_path / "o2"), "--force"]) in (0, 2)
    cfg = tmp_path / "bad.yaml"
    cfg.write_text("thresholds:\n  dormant_days: -5\n")
    assert main(["run", "--data", str(data), "--out", str(tmp_path / "o3"), "--config", str(cfg)]) == 1


def test_custom_config_changes_report(tmp_path):
    data, out = tmp_path / "d", tmp_path / "o"
    cfg = tmp_path / "c.yaml"
    cfg.write_text("thresholds:\n  dormant_days: 200\n")
    assert main(["all", "--data", str(data), "--out", str(out), "--config", str(cfg)]) == 0
    s = pd.read_csv(out / "csv/control_summary.csv").set_index("control_id")
    assert s.loc["T2", "exceptions"] < 34
