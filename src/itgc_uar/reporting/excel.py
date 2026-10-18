"""Excel audit workpaper (openpyxl). Every sheet is generated from the run - nothing is typed in by hand."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from itgc_uar import __version__
from itgc_uar.config import CONTROL_IDS
from itgc_uar.controls import CONTROLS
from itgc_uar.engine import RunResult

NAVY = "104281"
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
SEV_FILL = {"High": "F4B6B0", "Medium": "FBE3A6", "Low": "D5E8D4"}
STATUS_FILL = {"Not effective": "F4B6B0", "Effective": "D5E8D4", "MATCH": "D5E8D4", "MISMATCH": "F4B6B0",
               "FAIL": "F4B6B0", "PASS": "D5E8D4"}
NICE = {"control_id": "Control ID", "control": "Control", "control_statement": "Control statement",
        "population_basis": "Population basis", "population": "Population tested", "exceptions": "Exceptions",
        "deviation_rate": "Deviation rate", "tolerable_rate": "Tolerable rate", "conclusion": "Conclusion",
        "high": "High", "medium": "Medium", "low": "Low", "p1": "P1", "p2": "P2", "p3": "P3", "p4": "P4",
        "avg_risk_score": "Avg risk score", "max_risk_score": "Max risk score", "owner": "Suggested owner",
        "reconciliation": "SQL vs pandas"}


def _style(ws, df: pd.DataFrame, widths: dict[str, int] | None = None, wrap: tuple[str, ...] = ()) -> None:
    for c in ws[1]:
        c.fill, c.font = HEADER_FILL, Font(bold=True, color="FFFFFF")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(df.columns, start=1):
        sample = [len(str(col))] + [len(str(v)) for v in df[col].head(200)]
        w = (widths or {}).get(col, min(60, max(10, max(sample) + 2)))
        ws.column_dimensions[get_column_letter(i)].width = w
        for cell in ws.iter_cols(min_col=i, max_col=i, min_row=2):
            for x in cell:
                if col in wrap:
                    x.alignment = Alignment(wrap_text=True, vertical="top")
                v = str(x.value)
                if col in ("severity", "conclusion", "Severity", "Conclusion", "reconciliation", "SQL vs pandas",
                           "Status") and v in (SEV_FILL | STATUS_FILL):
                    x.fill = PatternFill("solid", fgColor=(SEV_FILL | STATUS_FILL)[v])


def _write(writer, df: pd.DataFrame, name: str, **kw) -> None:
    df.to_excel(writer, sheet_name=name, index=False)
    _style(writer.sheets[name], df, **kw)


def write_excel(result: RunResult, out_dir: Path, validation: pd.DataFrame | None, meta: dict) -> Path:
    path = out_dir / "ITGC_UAR_Exception_Report.xlsx"
    out_dir.mkdir(parents=True, exist_ok=True)
    cfg = result.cfg
    s = result.summary
    reg = result.register
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        cover = pd.DataFrame({"Item": [
            "DATA STATUS", "Report", "Review date", "Controls tested", "Total exceptions", "High severity",
            "P1 priority", "SQL vs pandas reconciliation", "Data-quality errors / warnings", "Config fingerprint",
            "Tool version", "Dataset", "How to read this file"], "Value": [
            "ALL DATA IN THIS WORKBOOK IS SIMULATED - no real organisation, employee or system.",
            "ITGC User Access Review - control testing exception report", cfg.review_date.isoformat(),
            len(CONTROL_IDS), len(reg), int((reg["severity"] == "High").sum()), int((reg["priority"] == "P1").sum()),
            "ALL MATCH" if result.reconciled else "MISMATCH - see 'Reconciliation' sheet",
            f"{result.dq.errors} / {result.dq.warnings}", cfg.fingerprint(), __version__,
            f"{meta.get('profile', 'custom')} (seed {meta.get('seed', 'n/a')})",
            "Summary = results per control. Exception Register = every finding with evidence. Remediation Plan = what "
            "to fix first. T1-T6 = raw test output. Reconciliation = proof the SQL and pandas results agree."]})
        cover.to_excel(xw, sheet_name="Cover", index=False)
        ws = xw.sheets["Cover"]
        _style(ws, cover, widths={"Item": 34, "Value": 110}, wrap=("Value",))
        ws["B2"].font = Font(bold=True, color="B00020")

        sdf = s.rename(columns=NICE)
        _write(xw, sdf, "Summary", wrap=("Control statement",), widths={"Control statement": 48, "Population basis": 34})
        ws = xw.sheets["Summary"]
        for col in ("G", "H"):
            for c in ws[col][1:]:
                c.number_format = "0.0%"

        _write(xw, reg, "Exception Register", wrap=("evidence", "recommended_action", "score_factors"),
               widths={"evidence": 70, "recommended_action": 55, "score_factors": 50, "exception_type": 36})

        plan = reg.sort_values(["priority", "risk_score", "exception_id"], ascending=[True, False, True])[
            ["priority", "due_date", "exception_id", "control_id", "exception_type", "severity", "risk_score",
             "department", "username", "system", "owner", "recommended_action"]]
        _write(xw, plan, "Remediation Plan", wrap=("recommended_action",), widths={"recommended_action": 60})

        for cid in CONTROL_IDS:
            d = result.details[cid]
            _write(xw, d, cid, wrap=("evidence", "recommended_action", "score_factors"),
                   widths={"evidence": 70, "recommended_action": 55, "score_factors": 50})
        dept = result.dept_matrix.reset_index().rename(columns={"department": "Department"})
        _write(xw, dept, "By Department")
        _write(xw, result.dq.to_frame(), "Data Quality", wrap=("Description", "Detail"), widths={"Description": 70})
        rec = pd.DataFrame([{
            "Control ID": r.control_id, "Control": CONTROLS[r.control_id].name, "SQL exceptions": r.sql_count,
            "pandas exceptions": r.pandas_count, "Keys matched": r.matched_keys, "Only in SQL": len(r.only_sql),
            "Only in pandas": len(r.only_pandas), "Attribute differences": len(r.attribute_mismatches),
            "SQL population": r.sql_population, "pandas population": r.pandas_population,
            "Status": r.status, "Examples": r.examples()} for r in result.recon.values()])
        _write(xw, rec, "Reconciliation", widths={"Examples": 80})
        if validation is not None:
            v = validation.copy()
            v.insert(0, "note", "SIMULATED data only: compares detections with what the generator planted.")
            _write(xw, v.drop(columns=["note"]), "Synthetic Validation")
        cfgdf = pd.DataFrame([{"Setting": k, "Value": str(v)} for k, v in cfg.as_dict().items()])
        _write(xw, cfgdf, "Config", widths={"Value": 90})
    return path
